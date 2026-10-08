"""Excel-звіт: порівняння, журнал змін, ціни по днях, пари на перевірку, підсумок."""

from __future__ import annotations

from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
HEAD_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
WHITE_BOLD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
RED_FONT = Font(name="Arial", size=10, bold=True, color="B8281C")
GREEN_FONT = Font(name="Arial", size=10, color="2E7D4F")
LINK_FONT = Font(name="Arial", size=10, color="1F5FBF", underline="single")
NOTE_FONT = Font(name="Arial", size=10, italic=True, color="5E6D62")
HEAD_FILL = PatternFill("solid", fgColor="2F4A38")
RED_ROW = PatternFill("solid", fgColor="FBE3DF")
RED_CELL = PatternFill("solid", fgColor="B8281C")
ORANGE_ROW = PatternFill("solid", fgColor="FCEFD9")
ORANGE_CELL = PatternFill("solid", fgColor="9A5B00")
GREEN_CELL = PatternFill("solid", fgColor="DCEFE2")
MONEY = "#,##0.00"
PCT = "0.0%"
STATUS = {"red": "Потребує уваги", "orange": "Ваша ціна вища", "ok": "В нормі"}
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def _header(ws, headers, row=1):
    for i, (title, width) in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=title)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 32


def _put(ws, row, col, value, fmt=None, font=FONT, align=TOP):
    c = ws.cell(row=row, column=col, value=value)
    c.font, c.alignment = font, align
    if fmt:
        c.number_format = fmt
    return c


def _link(ws, row, col, url):
    if not url:
        return
    c = _put(ws, row, col, "відкрити", font=LINK_FONT)
    c.hyperlink = url


def sup_list(ctx):
    return ctx.get("suppliers") or [{"id": "", "label": ctx["labels"]["sup"]}]


def cell_of(row, sup_id):
    for c in row["cells"]:
        if c["sup"] == sup_id:
            return c
    return None


def sheet_compare(wb, ctx):
    ws = wb.active
    ws.title = "Порівняння"
    sups = sup_list(ctx)
    headers = [("Статус", 17), ("Що сталось", 54), ("Мій ID", 11), ("Мій товар", 46),
               ("Моя ціна, грн", 13), ("У мене в наявності", 12)]
    for sp in sups:
        headers += [(f"{sp['label']}: ціна, грн", 14), (f"{sp['label']}: наявність", 12)]
    first_sup = 7
    ref_col = first_sup + len(sups) * 2
    headers += [("Орієнтир", 18), ("Ціна орієнтира, грн", 14), ("Різниця, грн", 12), ("Різниця, %", 11),
                ("Зміна в орієнтира сьогодні, грн", 16), ("Змін в орієнтира за 30 днів", 14),
                ("Розбіжність, днів", 12), ("Як зіставлено", 16), ("Мій товар", 11)]
    _header(ws, headers)
    last_col = get_column_letter(len(headers))
    price_col = get_column_letter(5)
    refp = get_column_letter(ref_col + 1)
    diff = get_column_letter(ref_col + 2)

    for i, r in enumerate(ctx["rows"], start=2):
        _put(ws, i, 1, STATUS[r["sev"]])
        _put(ws, i, 2, r["what"] or "", align=WRAP_TOP)
        _put(ws, i, 3, r["mid"])
        _put(ws, i, 4, r["mname"], align=WRAP_TOP)
        _put(ws, i, 5, r["mp"], MONEY)
        _put(ws, i, 6, "так" if r["mavail"] else "ні", font=FONT if r["mavail"] else RED_FONT)
        for n, sp in enumerate(sups):
            c = cell_of(r, sp["id"])
            col = first_sup + n * 2
            _put(ws, i, col, c["sp"] if c else None, MONEY,
                 font=GREEN_FONT if c and r["ref"] == sp["id"] and len(sups) > 1 else FONT)
            _put(ws, i, col + 1, ("так" if c["savail"] else "ні") if c else "",
                 font=RED_FONT if c and not c["savail"] else FONT)
        _put(ws, i, ref_col, next((sp["label"] for sp in sups if sp["id"] == r["ref"]), ""))
        _put(ws, i, ref_col + 1, r["sp"], MONEY)
        _put(ws, i, ref_col + 2, f'=IF({refp}{i}="","",{price_col}{i}-{refp}{i})', MONEY)
        _put(ws, i, ref_col + 3, f'=IF(OR({refp}{i}="",{refp}{i}=0),"",{diff}{i}/{refp}{i})', PCT)
        _put(ws, i, ref_col + 4, r["sch"] or None, MONEY, font=RED_FONT if r["sch"] else FONT)
        _put(ws, i, ref_col + 5, r["c30"])
        _put(ws, i, ref_col + 6, r["days"])
        _put(ws, i, ref_col + 7, r["how"])
        _link(ws, i, ref_col + 8, r["murl"])

        if r["sev"] in ("red", "orange"):
            row_fill = RED_ROW if r["sev"] == "red" else ORANGE_ROW
            for col in range(2, len(headers) + 1):
                ws.cell(row=i, column=col).fill = row_fill
            status = ws.cell(row=i, column=1)
            status.fill = RED_CELL if r["sev"] == "red" else ORANGE_CELL
            status.font = WHITE_BOLD
        else:
            ws.cell(row=i, column=1).fill = GREEN_CELL

    last_row = max(len(ctx["rows"]), 1) + 1
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"
    return last_row, get_column_letter(ref_col + 4)


def sheet_log(wb, ctx):
    ws = wb.create_sheet("Журнал змін")
    headers = [("Дата", 12), ("Хто змінив", 18), ("Мій ID", 11), ("ID у постачальника", 13), ("Товар", 56),
               ("Було, грн", 12), ("Стало, грн", 12), ("Зміна, грн", 12), ("Зміна, %", 10)]
    _header(ws, headers)
    for i, e in enumerate(ctx["log"], start=2):
        supplier = e["side"] == "постачальник"
        _put(ws, i, 1, date.fromisoformat(e["d"]), "DD.MM.YYYY")
        _put(ws, i, 2, e.get("supLabel", "Постачальник") if supplier else "Ваша ціна")
        _put(ws, i, 3, e["mid"])
        _put(ws, i, 4, e["sid"])
        _put(ws, i, 5, e["name"], align=WRAP_TOP)
        _put(ws, i, 6, e["old"], MONEY)
        _put(ws, i, 7, e["new"], MONEY)
        _put(ws, i, 8, f"=G{i}-F{i}", MONEY)
        _put(ws, i, 9, f'=IF(F{i}=0,"",H{i}/F{i})', PCT)
        if supplier:
            for col in range(1, len(headers) + 1):
                ws.cell(row=i, column=col).fill = RED_ROW
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{max(len(ctx['log']), 1) + 1}"


def sheet_by_day(wb, title, ctx, hist, hist_dates, today, key, changed_fill):
    """Ціни по днях: окремий рядок на кожну пару «мій товар плюс постачальник»."""
    ws = wb.create_sheet(title)
    sups = sup_list(ctx)
    dates = (hist_dates + [today])[-31:]
    headers = [("Постачальник", 18), ("Мій ID", 11), ("ID у постачальника", 13), ("Товар", 46)]
    headers += [(d.strftime("%d.%m"), 10) for d in dates]
    _header(ws, headers)
    field = "sp" if key == "sp" else "mp"
    i = 1
    for r in ctx["rows"]:
        for c in r["cells"]:
            i += 1
            label = next((sp["label"] for sp in sups if sp["id"] == c["sup"]), c["sup"])
            by_date = {e["d"]: e[field] for e in hist.get(tuple(c["key"]), [])}
            by_date[today] = c["sp"] if field == "sp" else r["mp"]
            _put(ws, i, 1, label)
            _put(ws, i, 2, r["mid"])
            _put(ws, i, 3, c["sid"])
            _put(ws, i, 4, r["mname"])
            prev = None
            for j, d in enumerate(dates, start=5):
                v = by_date.get(d)
                cell = _put(ws, i, j, v, MONEY)
                if v is not None and prev is not None and abs(v - prev) >= 0.01:
                    cell.fill = changed_fill
                if v is not None:
                    prev = v
    ws.freeze_panes = "E2"


def sheet_suggestions(wb, ctx):
    ws = wb.create_sheet("Пари на перевірку")
    ws.merge_cells("A1:I1")
    note = ws["A1"]
    note.value = ("Скрипт не наважився зіставити ці товари сам. Щоб підтвердити пару, відкрийте mapping/suggestions.csv "
                  "на GitHub і впишіть «так» у колонку «рішення», щоб відхилити, «ні».")
    note.font, note.alignment = NOTE_FONT, Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 34
    headers = [("Схожість", 10), ("Постачальник", 18), ("Мій ID", 11), ("ID у постачальника", 13), ("Мій товар", 46),
               ("Товар постачальника", 46), ("Моя ціна, грн", 13), ("Ціна постачальника, грн", 14), ("Причина", 34)]
    _header(ws, headers, row=2)
    for i, s in enumerate(ctx["sugg"], start=3):
        _put(ws, i, 1, s["score"])
        _put(ws, i, 2, s.get("supLabel", ""))
        _put(ws, i, 3, s["mid"])
        _put(ws, i, 4, s["sid"])
        _put(ws, i, 5, s["mname"], align=WRAP_TOP)
        _put(ws, i, 6, s["sname"], align=WRAP_TOP)
        _put(ws, i, 7, s["mp"], MONEY)
        _put(ws, i, 8, s["sp"], MONEY)
        _put(ws, i, 9, s["reason"], align=WRAP_TOP)
    ws.freeze_panes = "A3"


def sheet_summary(wb, ctx, last_row, sch_col):
    ws = wb.create_sheet("Підсумок")
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 44
    ref = f"'Порівняння'!A2:A{last_row}"
    totals = ctx["stats"].get("sup_totals") or {}
    rows = [
        ("Звіт сформовано", ctx["generated"], ""),
        ("Спільних товарів", f"=COUNTA({ref})", ""),
        ("Потребують уваги", f'=COUNTIF({ref},"Потребує уваги")', "червоні рядки на аркуші «Порівняння»"),
        ("Ваша ціна вища за орієнтир", f'=COUNTIF({ref},"Ваша ціна вища")', ""),
        ("В нормі", f'=COUNTIF({ref},"В нормі")', ""),
        ("Постачальник змінив ціну сьогодні", f"=COUNT('Порівняння'!{sch_col}2:{sch_col}{last_row})", ""),
        ("Немає в наявності в жодного постачальника", ctx["counts"].get("nostock", 0), "товар продається, а купити ніде"),
        ("Товарів у вашому фіді", ctx["stats"]["my_total"], ""),
    ]
    for sp in sup_list(ctx):
        rows.append((f"Товарів у «{sp['label']}»", totals.get(sp["id"], ctx["stats"]["sup_total"]), ""))
    rows += [
        ("Пар на перевірку", len(ctx["sugg"]), "аркуш «Пари на перевірку»"),
        ("Як зіставлено", ", ".join(f"{k}: {v}" for k, v in ctx["stats"]["by_method"].items()) or "нічого", ""),
        ("", "", ""),
        ("Пороги", "", "значення з config.json у репозиторії"),
        ("Допуск на округлення, грн", ctx["thresholds"]["tolerance_uah"], "різниця в межах допуску не підсвічується"),
        ("Мінімальна націнка до ціни постачальника, %", ctx["thresholds"]["min_markup_pct"], "нижче цього рядок червоний"),
        ("Максимальна націнка до ціни постачальника, %", ctx["thresholds"]["max_markup_pct"], "вище цього рядок помаранчевий"),
    ]
    _header(ws, [("Показник", 48), ("Значення", 22), ("Пояснення", 44)])
    for i, (a, b, c) in enumerate(rows, start=2):
        _put(ws, i, 1, a, font=BOLD if a == "Пороги" else FONT)
        _put(ws, i, 2, b)
        _put(ws, i, 3, c, font=NOTE_FONT)


def write_xlsx(path, ctx, hist, hist_dates, today):
    wb = Workbook()
    last_row, sch_col = sheet_compare(wb, ctx)
    sheet_log(wb, ctx)
    sheet_by_day(wb, "Ціни постачальників по днях", ctx, hist, hist_dates, today, "sp", RED_ROW)
    sheet_by_day(wb, "Ваші ціни по днях", ctx, hist, hist_dates, today, "mp", GREEN_CELL)
    sheet_suggestions(wb, ctx)
    sheet_summary(wb, ctx, last_row, sch_col)
    wb.save(path)
