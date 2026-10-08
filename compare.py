#!/usr/bin/env python3
"""
Щоденне порівняння цін: мій фід проти фідів і прайсів постачальників.

1. Завантажує мій фід і джерела всіх постачальників із config.json
   (YML, Prom XML, CSV або Excel, за посиланням або файлом у репозиторії).
2. Шукає спільні товари: ручні пари з mapping/mapping.csv, код / артикул / штрихкод,
   однакова назва, дуже схожа назва. Сумнівні пари йдуть у mapping/suggestions.csv.
3. Зберігає знімок цін у data/snapshots/РРРР-ММ-ДД.csv і журнал змін у data/changes.csv.
4. Будує звіти: docs/index.html, report.xlsx, README.md.

Змінні середовища:
  MY_FEED_URL          мій фід (тримати в секретах GitHub)
  SUPPLIER_FEED_URL    посилання для першого постачальника, якщо його не можна тримати в config.json
  SUPPLIER_URLS        посилання для решти: рядки виду «id = https://...», по одному на постачальника
  TELEGRAM_BOT_TOKEN   необов'язково
  TELEGRAM_CHAT_ID     необов'язково
  FEED_PROXY           необов'язково, проксі тільки для завантаження фідів (http://user:pass@host:port)
  DASHBOARD_PASSWORD   необов'язково, пароль веб-версії; тоді збирається зашифрована сторінка site/index.html
  SITE_REPO            необов'язково, публічний репозиторій для веб-версії у форматі власник/назва
  SKIP_TELEGRAM        необов'язково, будь-яке значення вимикає сповіщення
  RUN_DATE             необов'язково, РРРР-ММ-ДД, щоб перезапустити за певну дату
"""

from __future__ import annotations

import base64
import csv
import gzip
import hashlib
import html as htmllib
import io
import json
import os
import re
import sys
import time
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import quote
from zoneinfo import ZoneInfo

import requests
from lxml import etree
from rapidfuzz import fuzz

ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / "data"
SNAP_DIR = DATA_DIR / "snapshots"
CHANGES_CSV = DATA_DIR / "changes.csv"
LAST_RUN = DATA_DIR / "last_run.json"
MAP_DIR = ROOT / "mapping"
MAPPING_CSV = MAP_DIR / "mapping.csv"
SUGG_CSV = MAP_DIR / "suggestions.csv"
DOCS_DIR = ROOT / "docs"
SITE_DIR = ROOT / "site"
TEMPLATE = ROOT / "template.html"
SITE_TEMPLATE = ROOT / "site_template.html"
PBKDF2_ITER = 600_000
KYIV = ZoneInfo("Europe/Kyiv")
UA = "Mozilla/5.0 (compatible; PriceWatch/1.0)"

SIDE_SUP = "постачальник"
SIDE_MY = "моя"
YES = {"так", "да", "yes", "y", "1", "+", "ok", "ок"}
NO = {"ні", "нет", "no", "n", "0", "-"}

DEFAULT_CONFIG = {
    "my_label": "Мій магазин",
    "suppliers": [],
    "tolerance_uah": 1,
    "min_markup_pct": 0,
    "max_markup_pct": 0,
    "code_match_min_name_score": 45,
    "match_by_supplier_offer_id": True,
    "fuzzy_auto_accept": 90,
    "fuzzy_suggest": 72,
    "report_history_days": 90,
    "log_days": 60,
    "telegram_only_if_attention": True,
    "telegram_attach_dashboard": True,
    "github_pages": False,
}


# ---------------------------------------------------------------- утиліти

def load_config() -> dict:
    cfg = dict(DEFAULT_CONFIG)
    path = ROOT / "config.json"
    if path.exists():
        cfg.update(json.loads(path.read_text(encoding="utf-8")))
    return cfg


def load_suppliers(cfg: dict) -> list[dict]:
    """Постачальники з config.json. Старий формат з одним постачальником теж читається."""
    raw = cfg.get("suppliers") or []
    if not raw:
        raw = [{"id": "supplier", "label": cfg.get("supplier_label") or "Постачальник",
                "url": cfg.get("supplier_feed_url") or ""}]
    secret = env_urls()
    out, seen = [], set()
    for i, sp in enumerate(raw):
        sid = re.sub(r"[^a-z0-9_-]", "", str(sp.get("id") or f"s{i + 1}").strip().lower()) or f"s{i + 1}"
        while sid in seen:
            sid += "x"
        seen.add(sid)
        src = secret.get(sid) or str(sp.get("url") or sp.get("file") or "").strip()
        if i == 0 and os.environ.get("SUPPLIER_FEED_URL", "").strip():
            src = os.environ["SUPPLIER_FEED_URL"].strip()
        out.append({"id": sid, "label": str(sp.get("label") or sid), "src": src,
                    "columns": sp.get("columns") or {}, "format": str(sp.get("format") or "").lower(),
                    "scraper": str(sp.get("scraper") or "").strip().lower(), "options": sp.get("options") or {}})
    return out


SCRAPED_DIR = ROOT / "data" / "scraped"
SCRAPERS = {"plants_club": "scrapers.plants_club"}


def run_scraper(sp: dict) -> dict:
    """Збирає каталог сайту без фіду і зберігає зібране в data/scraped/<id>.csv для перевірки."""
    import importlib

    module_name = SCRAPERS.get(sp["scraper"])
    if not module_name:
        raise FeedError(f"Невідомий збирач «{sp['scraper']}» у постачальника «{sp['label']}». "
                        f"Доступні: {', '.join(sorted(SCRAPERS))}.")
    module = importlib.import_module(module_name)
    print(f"  Збираю каталог «{sp['label']}» з сайту...")
    try:
        rows = module.scrape(sp["options"], log=print)
    except module.ScrapeError as exc:
        raise FeedError(f"Не вдалося зібрати каталог «{sp['label']}»: {exc}.") from exc

    SCRAPED_DIR.mkdir(parents=True, exist_ok=True)
    with (SCRAPED_DIR / f"{sp['id']}.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["id", "артикул", "назва", "виробник", "категорія", "ціна", "в_наявності", "залишок", "посилання"])
        for r in sorted(rows, key=lambda r: r["name"]):
            w.writerow([r["id"], r["sku"], r["name"], r["brand"], r["category"], fmt_num(r["price"]),
                        "так" if r["available"] else "ні", "" if r["stock"] is None else fmt_num(r["stock"]), r["url"]])

    offers = {}
    for r in rows:
        codes = {c for c in [norm_code(r["sku"])] if field_code_ok(c)}
        offers[r["id"]] = make_offer(r["id"], r["name"], r["price"], r["url"], r["available"], codes)
    return offers


def fmt_num(v) -> str:
    if v is None:
        return ""
    return f"{v:.2f}".rstrip("0").rstrip(".")


def pct(new, old):
    return round((new - old) / old * 100, 1) if old else None


def plural(n: int, one: str, few: str, many: str) -> str:
    a = abs(n) % 100
    b = a % 10
    if 10 < a < 20:
        return many
    if b == 1:
        return one
    if 2 <= b <= 4:
        return few
    return many


def read_text_any(path: Path) -> str:
    """CSV могли зберегти в Excel у Windows-1251, тому пробуємо обидва кодування."""
    raw = path.read_bytes()
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1251", errors="replace")


def read_rows(path: Path) -> list[dict]:
    if not path.exists():
        return []
    text = read_text_any(path)
    if not text.strip():
        return []
    head = text.splitlines()[0]
    delim = ";" if head.count(";") > head.count(",") else ","
    rows = []
    for r in csv.DictReader(io.StringIO(text), delimiter=delim, restkey="_зайве"):
        row = {}
        for k, v in r.items():
            if isinstance(v, list):  # у рядку більше колонок, ніж у заголовку
                v = v[0] if v else ""
            row[(k or "").strip().lower()] = str(v or "").strip()
        rows.append(row)
    return rows


def pick(row: dict, *names: str) -> str:
    for n in names:
        if row.get(n):
            return row[n]
    return ""


# ---------------------------------------------------------------- нормалізація назв і кодів

CYR = "А-ЯЁЄІЇҐ"
HOMO = str.maketrans("АВЕКМНОРСТХІУ", "ABEKMHOPCTXIY")  # кирилиця, схожа на латиницю
TOKEN_RE = re.compile(rf"[0-9A-Z{CYR}]+(?:\.[0-9]+)?")
JOIN_RE = re.compile(rf"(?<=[0-9A-Z])[-/](?=[0-9A-Z{CYR}])|(?<=[0-9A-Z{CYR}])[-/](?=[0-9A-Z])")
UNIT_RE = re.compile(r"(?<=\d)(?=(?:ММ|СМ|МЛ|КГ|ШТ|ВТ|М|Л|Г|MM|CM|ML|KG|W|V|M|L|G)\b)")
PAREN_RE = re.compile(r"\(([^()]{3,40})\)")
UNIT_LIKE = re.compile(r"^\d+(?:[.,]\d+)?\s*[^\W\d_]{1,4}\.?$")
LONG_TOKEN_RE = re.compile(r"(?<![0-9A-Za-z])[0-9A-Za-z][0-9A-Za-z.\-/]{4,}")
UNIT_CODE = re.compile(r"\d+(?:W|V|MM|CM|ML|KG|AH|MAH|RPM|NM|KW|L|G|M)")


def _prep(text: str) -> str:
    s = (text or "").upper().replace("Ё", "Е")
    s = re.sub(r"[\"'`’ʼ«»]", "", s).replace(",", ".")
    s = JOIN_RE.sub("", s)
    return UNIT_RE.sub(" ", s)


def _fix_token(tok: str) -> str:
    if re.search(r"[0-9A-Z]", tok) and re.search(rf"[{CYR}]", tok):
        return tok.translate(HOMO)
    return tok


def name_tokens(text: str) -> list[str]:
    return [_fix_token(t) for t in TOKEN_RE.findall(_prep(text))]


def norm_code(text: str) -> str:
    return re.sub(rf"[^0-9A-Z{CYR}]", "", (text or "").upper().translate(HOMO))


def field_code_ok(c: str) -> bool:
    if len(c) < 4 or len(set(c)) == 1:
        return False
    return any(ch.isdigit() for ch in c) or len(c) >= 6


def codes_from_name(name: str) -> set[str]:
    out = set()
    for inner in PAREN_RE.findall(name):
        for part in re.split(r"[;,]", inner):
            part = part.strip()
            if not part or part.count(" ") > 1 or UNIT_LIKE.match(part):
                continue
            c = norm_code(part)
            if len(c) >= 5 and len(set(c)) > 1 and any(ch.isdigit() for ch in c):
                out.add(c)
    for tok in LONG_TOKEN_RE.findall(name):
        c = norm_code(tok)
        if len(c) >= 6 and sum(ch.isdigit() for ch in c) >= 3 and not UNIT_CODE.fullmatch(c):
            out.add(c)
    return out


# Одиниці зводимо до базових, щоб «0.5 л» і «500 мл» були одним значенням.
UNIT_BASE = {
    "МЛ": ("об'єм", 1), "Л": ("об'єм", 1000), "ML": ("об'єм", 1), "L": ("об'єм", 1000),
    "Г": ("вага", 1), "КГ": ("вага", 1000), "G": ("вага", 1), "KG": ("вага", 1000), "Т": ("вага", 1000000),
    "ММ": ("довжина", 1), "СМ": ("довжина", 10), "М": ("довжина", 1000), "MM": ("довжина", 1), "CM": ("довжина", 10),
    "ШТ": ("кількість", 1), "PCS": ("кількість", 1),
    "ВТ": ("потужність", 1), "КВТ": ("потужність", 1000), "W": ("потужність", 1), "KW": ("потужність", 1000),
    "V": ("напруга", 1),  # кирилична «в» не рахується: це прийменник, як у «5 в 1»
    "АГ": ("ємність", 1), "МАГ": ("ємність", 0.001), "AH": ("ємність", 1), "MAH": ("ємність", 0.001),
}
MEASURE_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(" + "|".join(sorted(UNIT_BASE, key=len, reverse=True)) + r")(?![0-9A-Z" + CYR + r"])")
LAT_WORD = re.compile(r"^[A-Z]{2,}$")
CYR_WORD = re.compile(rf"^[{CYR}]{{4,}}$")
# слова, що не змінюють товар: країна походження, маркетинг, стан пакування
NOISE_WORDS = {
    "ЯПОНІЯ", "НІМЕЧЧИНА", "ПОЛЬЩА", "ІТАЛІЯ", "КИТАЙ", "УКРАЇНА", "ФРАНЦІЯ", "ІСПАНІЯ", "ТУРЕЧЧИНА",
    "ЧЕХІЯ", "ШВЕЦІЯ", "ШВЕЙЦАРІЯ", "АВСТРІЯ", "УГОРЩИНА", "СЛОВАЧЧИНА", "БЕЛЬГІЯ", "ГОЛЛАНДІЯ",
    "НІДЕРЛАНДИ", "ІЗРАЇЛЬ", "КОРЕЯ", "ТАЙВАНЬ", "ІНДІЯ", "США", "КАНАДА", "БРАЗИЛІЯ", "МОЛДОВА",
    "ОРИГІНАЛ", "ОРИГІНАЛЬНИЙ", "НОВИНКА", "НОВИЙ", "НОВА", "НОВЕ", "НОВІ", "АКЦІЯ", "ХІТ", "ТОП",
    "СУПЕР", "ФІРМОВИЙ", "ЯКІСНИЙ", "КРАЩИЙ", "ПОПУЛЯРНИЙ", "РОЗПРОДАЖ", "ЗНИЖКА",
    "ВИРОБНИЦТВО", "ВИРОБНИК", "ГАРАНТІЯ", "ДОСТАВКА", "НАЯВНОСТІ", "СКЛАДУ", "ПОДАРУНОК",
    "УПАКОВКА", "УПАКОВЦІ", "ПАКОВАННЯ", "ФАСОВКА", "ФАСОВАНИЙ", "ВАГОВИЙ", "ЗАВОДСЬКА", "ЗАВОДСЬКИЙ",
}


def measures(name: str) -> frozenset:
    """Числа з одиницями: «5 л» -> ('об'єм', 5000). Порожньо, якщо в назві немає жодної."""
    out = set()
    for num, unit in MEASURE_RE.findall(_prep(name)):
        dim, mult = UNIT_BASE[unit]
        out.add((dim, round(float(num) * mult, 3)))
    return frozenset(out)


def measure_conflict(a: frozenset, b: frozenset) -> bool:
    """Конфлікт, якщо обидві назви мають розміри і жодна не є доповненням іншої."""
    if not a or not b or a == b:
        return False
    return not (a <= b or b <= a)


def split_words(tokens: list[str]) -> tuple[frozenset, frozenset]:
    """Латинські слова (найчастіше бренд) і значущі кириличні слова."""
    lat = {t for t in tokens if LAT_WORD.fullmatch(t) and t not in UNIT_BASE}
    cyr = {t for t in tokens if CYR_WORD.fullmatch(t) and t not in NOISE_WORDS}
    return frozenset(lat), frozenset(cyr)


def brand_diff(a: frozenset, b: frozenset) -> frozenset:
    """Бренди, наявні лише з одного боку. «B/S XL» і «B/S-XL» дають BS, XL і BSXL: це те саме."""
    out = set()
    for x, other in ((a - b, b), (b - a, a)):
        for t in x:
            if not any(t in o or o in t for o in other):
                out.add(t)
    return frozenset(out)


def pair_conflict(m: "Offer", s: "Offer") -> str:
    """Причина, з якої пару не можна приймати автоматично."""
    if measure_conflict(m.measures, s.measures):
        return "різні об'єми або розміри"
    if m.brands and s.brands and m.brands.isdisjoint(s.brands) and brand_diff(m.brands, s.brands):
        return "різні бренди"
    if m.name_codes and s.name_codes and m.name_codes.isdisjoint(s.name_codes):
        return "різні коди виробника"
    return ""


def name_score(a: str, b: str) -> float:
    if not a or not b:
        return 0.0
    return round((fuzz.token_set_ratio(a, b) + fuzz.token_sort_ratio(a, b)) / 2, 1)


# ---------------------------------------------------------------- фіди

@dataclass
class Offer:
    id: str
    name: str
    price: float | None
    url: str = ""
    available: bool = True
    field_codes: set = field(default_factory=set)
    name_codes: set = field(default_factory=set)
    measures: frozenset = frozenset()
    brands: frozenset = frozenset()
    words: frozenset = frozenset()
    codes: set = field(default_factory=set)
    norm: str = ""
    model: frozenset = frozenset()


def make_offer(oid, name, price, url, available, field_codes) -> Offer:
    name = re.sub(r"\s+", " ", name or "").strip()
    toks = name_tokens(name)
    name_codes = codes_from_name(name)
    code_like = {c.lower() for c in name_codes}
    brands, words = split_words(toks)
    return Offer(
        id=oid, name=name, price=price, url=url, available=available,
        field_codes=set(field_codes), name_codes=name_codes,
        measures=measures(name), brands=brands, words=words,
        codes=set(field_codes) | name_codes,
        norm=" ".join(toks).lower(),
        # цифрові токени моделі без артикулів: «GX-297 (7980569366)» і «GX-297» мають однакову модель
        model=frozenset(t.lower() for t in toks if any(ch.isdigit() for ch in t) and t.lower() not in code_like),
    )


def env_urls() -> dict:
    """Секрет SUPPLIER_URLS: рядки виду «id = https://...» для постачальників з закритим фідом."""
    out = {}
    for line in os.environ.get("SUPPLIER_URLS", "").splitlines():
        key, sep, val = line.partition("=")
        if sep and val.strip():
            out[key.strip().lower()] = val.strip()
    return out


class FeedError(RuntimeError):
    """Джерело не прочиталось. Для постачальника це попередження, для свого фіду зупинка."""


def load_feed(src: str, label: str) -> bytes:
    if not src:
        raise FeedError(f"Не задано посилання на фід «{label}». Додайте секрет у Settings > Secrets and variables > Actions.")
    if not src.lower().startswith(("http://", "https://")):
        path = src if Path(src).is_absolute() else ROOT / src
        if Path(path).exists():
            return Path(path).read_bytes()
        raise FeedError(f"Файл «{label}» не знайдено: {src}. Покладіть його в репозиторій і вкажіть шлях у config.json.")
    src = normalize_source(src)
    err = ""
    proxy = os.environ.get("FEED_PROXY", "").strip()
    proxies = {"http": proxy, "https": proxy} if proxy else None
    for attempt in range(1, 4):
        try:
            r = requests.get(src, timeout=(20, 240), proxies=proxies,
                             headers={"User-Agent": UA, "Accept": "application/xml,text/xml,*/*"})
            if r.status_code != 200:
                raise RuntimeError(f"сервер відповів HTTP {r.status_code}")
            head = r.content[:4000].lstrip().lower()
            if head[:9] == b"<!doctype" or head[:5] == b"<html":
                # це не тимчасовий збій, повторювати немає сенсу
                raise FeedError(f"«{label}»: посилання відкриває веб-сторінку, а не файл. "
                                f"Відкрийте доступ «усі, хто має посилання» або перевірте адресу.")
            if is_xml(src) and sniff_format(r.content, src) == "xml" \
                    and b"<offer" not in r.content and b"<item" not in r.content:
                raise RuntimeError("у відповіді немає товарів")
            return r.content
        except FeedError:
            raise
        except Exception as exc:  # noqa: BLE001
            err = str(exc).replace(src, "***")
            if proxy:
                err = err.replace(proxy, "***")
            print(f"  «{label}»: спроба {attempt} не вдалася: {err}")
            if attempt < 3:
                time.sleep(15 * attempt)
    raise FeedError(f"Не вдалося завантажити фід «{label}»: {err}")


STRING = etree.XPath("string()")
CODE_TAGS = {"vendorcode", "vendor_code", "sku", "article", "artikul", "barcode", "ean", "gtin", "mpn", "code"}
WANTED_TAGS = CODE_TAGS | {"id", "name", "name_ua", "model", "typeprefix", "vendor", "price", "url", "available",
                           "presence", "in_stock", "stock", "quantity_in_stock", "stock_quantity", "quantity"}
CODE_PARAM = re.compile(r"артикул|штрих|barcode|ean|gtin|sku|mpn|код", re.I)
CODE_PARAM_SKIP = re.compile(r"уктзед|зед|митн|hs\b", re.I)
FALSE_WORDS = {"false", "0", "no", "ні", "нет", "not_available", "out_of_stock", "outofstock", "unavailable",
               "немає", "нема", "немає в наявності", "відсутній", "відсутня", "отсутствует", "нет в наличии",
               "під замовлення", "под заказ", "закінчився", "закончился", "очікується", "ожидается", "-", "—"}


def parse_price(text):
    if text is None:
        return None
    s = re.sub(r"[^\d.,]", "", str(text)).replace(",", ".")
    if s.count(".") > 1:
        head, _, tail = s.rpartition(".")
        s = head.replace(".", "") + "." + tail
    try:
        v = round(float(s), 2)
    except ValueError:
        return None
    return v if v > 0 else None


TABLE_COLS = {
    "id": ("id", "ід", "код товару", "артикул", "код", "sku", "vendorcode", "vendor_code", "номенклатура"),
    "name": ("назва", "найменування", "товар", "номенклатура", "name", "product", "опис"),
    "price": ("ціна", "цена", "price", "вартість", "роздріб", "розница", "opt", "опт"),
    "avail": ("наявн", "залишок", "остат", "склад", "stock", "available", "кількість", "количество", "quantity"),
    "code": ("артикул", "штрих", "barcode", "ean", "sku", "mpn", "vendorcode", "vendor_code", "код виробника"),
    "url": ("посилання", "url", "link", "сторінка"),
}


CLOUD_RULES = [
    # Google Таблиця: беремо аркуш як xlsx
    (re.compile(r"^https://docs\.google\.com/spreadsheets/d/([\w-]+)"),
     lambda m, u: f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=xlsx"
                  + (f"&gid={g.group(1)}" if (g := re.search(r"[#&?]gid=(\d+)", u)) else "")),
    # файл на Google Диску
    (re.compile(r"^https://drive\.google\.com/file/d/([\w-]+)"),
     lambda m, u: f"https://drive.google.com/uc?export=download&confirm=t&id={m.group(1)}"),
    (re.compile(r"^https://drive\.google\.com/open\?id=([\w-]+)"),
     lambda m, u: f"https://drive.google.com/uc?export=download&confirm=t&id={m.group(1)}"),
    (re.compile(r"^https://drive\.google\.com/uc\?"),
     lambda m, u: u if "confirm=" in u else u + "&confirm=t"),
    # Dropbox
    (re.compile(r"^https://(?:www\.)?dropbox\.com/"),
     lambda m, u: re.sub(r"(?<=[?&])dl=0", "dl=1", u) if "dl=0" in u
                  else (u if "dl=1" in u or "raw=1" in u else u + ("&" if "?" in u else "?") + "dl=1")),
    # OneDrive і SharePoint
    (re.compile(r"^https://[\w.-]*(?:1drv\.ms|onedrive\.live\.com|sharepoint\.com)/"),
     lambda m, u: u if "download=1" in u else u + ("&" if "?" in u else "?") + "download=1"),
]


def normalize_source(src: str) -> str:
    """Посилання на файл у хмарі перетворює на пряме завантаження."""
    for pattern, build in CLOUD_RULES:
        m = pattern.search(src)
        if m:
            return build(m, src)
    return src


def sniff_format(content: bytes, src: str) -> str:
    """Формат визначається за вмістом, бо в хмарних посиланнях немає розширення файлу."""
    head = content[:4000].lstrip()
    if head[:2] == b"PK":
        return "xlsx"
    if head[:1] == b"<" or b"<offer" in head or b"<item" in head or b"yml_catalog" in head:
        return "xml"
    if re.search(r"\.(xlsx|xlsm)(\?|$)", src.lower()) or "format=xlsx" in src.lower():
        return "xlsx"
    if re.search(r"\.(csv|tsv)(\?|$)", src.lower()) or "format=csv" in src.lower():
        return "csv"
    return "xml" if re.search(r"\.(xml|yml)(\?|$)", src.lower()) else "csv"


def is_xml(src: str) -> bool:
    return not re.search(r"\.(csv|tsv|xlsx|xlsm)(\?|$)", src.lower())


def table_rows(content: bytes, src: str, fmt: str = "") -> list[list]:
    if (fmt or sniff_format(content, src)) == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        return [list(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1251", errors="replace")
    sample = "\n".join(text.splitlines()[:20])
    delim = max(";,\t|", key=sample.count) if sample else ","
    return [r for r in csv.reader(io.StringIO(text), delimiter=delim)]


def map_columns(header: list) -> dict:
    """Знаходить потрібні колонки за назвами заголовків, у будь-якому порядку."""
    cells = [str(h or "").strip().lower() for h in header]
    found = {}
    for key, variants in TABLE_COLS.items():
        for i, h in enumerate(cells):
            if h and i not in found.values() and any(v in h for v in variants):
                found[key] = i
                break
    return found


def parse_table(content: bytes, label: str, src: str, cols_cfg: dict, fmt: str = "") -> dict[str, Offer]:
    """Прайс постачальника у форматі CSV або Excel. Заголовок шукається в перших 15 рядках."""
    try:
        rows = table_rows(content, src, fmt)
    except Exception as exc:  # noqa: BLE001
        raise FeedError(f"Не вдалося прочитати прайс «{label}»: {exc}. "
                        f"Перевірте, що посилання відкриває сам файл, а не сторінку перегляду.")
    head_at, cols = None, {}
    for i, row in enumerate(rows[:15]):
        found = map_columns(row)
        if "name" in found and "price" in found:
            head_at, cols = i, found
            break
    if cols_cfg:
        header = [str(h or "").strip().lower() for h in (rows[head_at] if head_at is not None else (rows[0] if rows else []))]
        for key, title in cols_cfg.items():
            t = str(title).strip().lower()
            if t in header:
                cols[key] = header.index(t)
                head_at = head_at if head_at is not None else 0
    if head_at is None or "name" not in cols or "price" not in cols:
        raise FeedError(f"У прайсі «{label}» не знайдено колонок з назвою і ціною. "
                        f"Впишіть їх у config.json у розділі columns цього постачальника.")

    offers: dict[str, Offer] = {}
    cell = lambda row, key: (str(row[cols[key]]).strip() if key in cols and cols[key] < len(row) and row[cols[key]] is not None else "")  # noqa: E731
    for n, row in enumerate(rows[head_at + 1:], 1):
        if not row or all(c in (None, "") for c in row):
            continue
        name = cell(row, "name")
        price = parse_price(cell(row, "price"))
        if not name or price is None:
            continue
        oid = cell(row, "id") or norm_code(cell(row, "code")) or f"r{n}"
        if oid in offers:
            oid = f"{oid}-{n}"
        raw_avail = cell(row, "avail")
        if raw_avail == "":
            available = True
        elif re.fullmatch(r"[\d.,\s]+", raw_avail):
            available = parse_price(raw_avail) is not None
        else:
            available = raw_avail.strip().lower() not in FALSE_WORDS
        codes = {c for part in re.split(r"[;,|]", cell(row, "code")) if field_code_ok(c := norm_code(part))}
        offers[oid] = make_offer(oid, name, price, cell(row, "url"), available, codes)
    return offers


def parse_source(content: bytes, label: str, src: str, cols_cfg: dict, fmt: str = "") -> dict[str, Offer]:
    fmt = (fmt or sniff_format(content, src)).lower()
    if fmt in ("xml", "yml"):
        return parse_feed(content, label)
    return parse_table(content, label, src, cols_cfg, fmt)


def parse_feed(content: bytes, label: str) -> dict[str, Offer]:
    parser = etree.XMLParser(recover=True, huge_tree=True, resolve_entities=False, no_network=True,
                             load_dtd=False, remove_comments=True, remove_pis=True)
    try:
        root = etree.fromstring(content, parser)
    except etree.XMLSyntaxError as exc:
        raise FeedError(f"Фід «{label}» не читається як XML: {exc}")
    if root is None:
        raise FeedError(f"Фід «{label}» порожній або пошкоджений")

    elements = list(root.iter("{*}offer")) or list(root.iter("{*}item"))
    offers: dict[str, Offer] = {}
    for el in elements:
        vals: dict[str, str] = {}
        params: list[tuple[str, str]] = []
        for ch in el:
            if not isinstance(ch.tag, str):
                continue
            tag = etree.QName(ch).localname.lower()
            if tag == "param":
                params.append((ch.get("name") or "", STRING(ch).strip()))
            elif tag in WANTED_TAGS and tag not in vals:
                vals[tag] = STRING(ch).strip()

        oid = (el.get("id") or vals.get("id") or "").strip()
        if not oid or oid in offers:
            continue

        name = vals.get("name_ua") or vals.get("name") or ""
        if not name and vals.get("model"):
            name = " ".join(x for x in (vals.get("typeprefix"), vals.get("vendor"), vals.get("model")) if x)

        available = True
        attr = el.get("available")
        if attr is not None:
            available = attr.strip().lower() not in FALSE_WORDS
        else:
            for key in ("available", "presence", "in_stock", "stock"):
                if key in vals:
                    available = vals[key].lower() not in FALSE_WORDS
                    break
            else:
                for key in ("quantity_in_stock", "stock_quantity", "quantity"):
                    if vals.get(key):
                        try:
                            available = float(vals[key].replace(",", ".")) > 0
                        except ValueError:
                            pass
                        break

        field_codes = set()
        raw_codes = [v for k, v in vals.items() if k in CODE_TAGS and v]
        raw_codes += [v for n, v in params if v and CODE_PARAM.search(n) and not CODE_PARAM_SKIP.search(n)]
        for raw in raw_codes:
            for part in re.split(r"[;,|]", raw):
                c = norm_code(part)
                if field_code_ok(c):
                    field_codes.add(c)

        offers[oid] = make_offer(oid, name, parse_price(vals.get("price")), vals.get("url", ""), available, field_codes)
    return offers


# ---------------------------------------------------------------- ручні правила

MAP_HEADER = ["мій_id", "постачальник", "id_постачальника", "правило", "коментар"]
SUGG_HEADER = ["рішення", "постачальник", "схожість", "мій_id", "id_постачальника", "мій_товар",
               "товар_постачальника", "моя_ціна", "ціна_постачальника", "причина"]


def rule_supplier(row: dict, default: str) -> str:
    """Порожня колонка «постачальник» означає першого зі списку: так читаються старі файли."""
    return (pick(row, "постачальник", "supplier") or default).strip().lower()


def load_rules(suppliers: list[dict]):
    """Для кожного постачальника: примусові пари, заборонені пари, повністю виключені товари."""
    known = {sp["id"] for sp in suppliers}
    default = suppliers[0]["id"]
    forced = {sp["id"]: {} for sp in suppliers}
    blocked = {sp["id"]: set() for sp in suppliers}
    excluded = {sp["id"]: set() for sp in suppliers}
    for r in read_rows(MAPPING_CSV):
        mid = pick(r, "мій_id", "my_id")
        sid = pick(r, "id_постачальника", "supplier_id")
        rule = (pick(r, "правило", "rule") or "так").lower()
        who = rule_supplier(r, default)
        if not mid:
            continue
        targets = sorted(known) if who in ("*", "усі", "все", "all") else [who]
        for t in targets:
            if t not in known:
                continue
            if rule in NO:
                if sid in ("", "*"):
                    excluded[t].add(mid)
                    forced[t].pop(mid, None)
                else:
                    blocked[t].add((mid, sid))
                    if forced[t].get(mid) == sid:
                        forced[t].pop(mid)
            elif sid and sid != "*":
                forced[t][mid] = sid
                blocked[t].discard((mid, sid))
                excluded[t].discard(mid)
    return forced, blocked, excluded


def map_rows(default: str) -> list[list[str]]:
    """Читає mapping.csv у єдиний формат: мій_id, постачальник, id_постачальника, правило, коментар."""
    out = []
    for r in read_rows(MAPPING_CSV):
        mid = pick(r, "мій_id", "my_id")
        if not mid:
            continue
        out.append([mid, rule_supplier(r, default), pick(r, "id_постачальника", "supplier_id"),
                    (pick(r, "правило", "rule") or "так").lower(), pick(r, "коментар", "comment")])
    return out


def absorb_decisions(today: date, default: str) -> int:
    """Переносить рішення «так / ні» з suggestions.csv у mapping.csv."""
    decided = []
    for r in read_rows(SUGG_CSV):
        d = pick(r, "рішення", "decision").lower()
        mid, sid = pick(r, "мій_id", "my_id"), pick(r, "id_постачальника", "supplier_id")
        who = rule_supplier(r, default)
        if not mid or not sid or not d:
            continue
        if d in YES:
            decided.append([mid, who, sid, "так", f"підтверджено {today:%d.%m.%Y}"])
        elif d in NO:
            decided.append([mid, who, sid, "ні", f"відхилено {today:%d.%m.%Y}"])
    existing = map_rows(default)
    seen = {tuple(r[:4]) for r in existing}
    fresh = [r for r in decided if tuple(r[:4]) not in seen]
    text = read_text_any(MAPPING_CSV) if MAPPING_CSV.exists() else ""
    old_head = text.splitlines()[0] if text.strip() else ""
    delim = ";" if old_head.count(";") > old_head.count(",") else ","
    cols = [h.strip().lower() for h in old_head.split(delim)]
    needs_upgrade = bool(old_head) and "постачальник" not in cols
    if not fresh and not needs_upgrade:
        return 0
    MAP_DIR.mkdir(parents=True, exist_ok=True)
    with MAPPING_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=delim, lineterminator="\n")
        w.writerow(MAP_HEADER)
        w.writerows(existing + fresh)
    if needs_upgrade:
        print("mapping.csv переведено у формат з колонкою «постачальник».")
    return len(fresh)


# ---------------------------------------------------------------- зіставлення

@dataclass
class Pair:
    m: Offer
    s: Offer
    how: str
    score: float
    sup: str = ""

    @property
    def key(self) -> tuple[str, str, str]:
        return (self.sup, self.m.id, self.s.id)


# слова, що означають інше виконання або супутній товар, а не опис того самого товару
VARIANT_STEMS = ("ЛІВ", "ПРАВ", "ДИТЯЧ", "МІНІ", "MINI", "MAXI", "НАБІР", "НАБОР", "КОМПЛЕКТ", "КОБУР", "ЧОХОЛ",
                 "СУМК", "ЛЕЗ", "ПРУЖИН", "ЗАПЧАСТ", "ЗМІНН", "РЕМКОМПЛ", "НАСАДК", "РУЧК", "ДЕРЖАК", "ТОЧИЛ",
                 "КАРТРИДЖ", "ФІЛЬТР", "ПІДСТАВК", "ТРИМАЧ", "АДАПТЕР", "ПЕРЕХІДН")


def model_ident(o: "Offer") -> set[str]:
    """Позначення моделі: цифро-літерні токени і коди з назви, у нижньому регістрі."""
    return set(o.model) | {c.lower() for c in o.name_codes}


def strong_ident(ident: set[str]) -> set[str]:
    return {t for t in ident if len(t) >= 4 and any(ch.isdigit() for ch in t) and any(ch.isalpha() for ch in t)}


def is_bundle(o: "Offer") -> bool:
    return "+" in o.name or any(w.startswith(("НАБІР", "НАБОР", "КОМПЛЕКТ")) for w in o.words)


def model_match_ok(m: "Offer", s: "Offer") -> bool:
    """Та сама модель того самого бренду, а зайве в назві постачальника лише опис."""
    mi, si = model_ident(m), model_ident(s)
    if not strong_ident(mi) or not mi <= si:
        return False
    # зайві позначення постачальника мають бути повтором наших, як «(АРС VS-8XZ)» до «VS-8XZ»
    if any(not any(t in x for t in mi) for x in si - mi):
        return False
    if not m.brands or brand_diff(m.brands, s.brands) or measure_conflict(m.measures, s.measures):
        return False
    if is_bundle(m) != is_bundle(s):
        return False
    if m.words - s.words:  # у нашій назві є слово, якого немає в постачальника
        return False
    return not any(w.startswith(VARIANT_STEMS) for w in s.words - m.words)


def match_offers(mine, sup, cfg, forced, blocked, excluded, sup_id=""):
    code_idx, name_idx, tok_idx = defaultdict(set), defaultdict(set), defaultdict(set)
    id_idx = {norm_code(s.id): s.id for s in sup.values()}
    for s in sup.values():
        for c in s.codes:
            code_idx[c].add(s.id)
        if s.norm:
            name_idx[s.norm].add(s.id)
            for t in set(s.norm.split()):
                if len(t) >= 3 or any(ch.isdigit() for ch in t):
                    tok_idx[t].add(s.id)
    df_limit = max(60, len(sup) // 20)
    ident_idx = defaultdict(set)
    for o in sup.values():
        for t in strong_ident(model_ident(o)):
            ident_idx[t].add(o.id)

    pairs, suggestions, warnings = [], [], []
    for m in mine.values():
        if m.id in excluded:
            continue
        if m.id in forced:
            sid = forced[m.id]
            if sid in sup:
                pairs.append(Pair(m, sup[sid], "вручну", 100.0, sup_id))
            else:
                warnings.append(f"Ручна пара {m.id} ↔ {sid}: товару постачальника {sid} немає у фіді.")
            continue

        fallback = None

        # 1. код, артикул, штрихкод
        cands = {sid for c in m.codes for sid in code_idx.get(c, ()) if (m.id, sid) not in blocked}
        method, min_score = "код", cfg["code_match_min_name_score"]
        if not cands and cfg.get("match_by_supplier_offer_id"):
            cands = {id_idx[c] for c in m.field_codes if c in id_idx and (m.id, id_idx[c]) not in blocked}
            method, min_score = "ID постачальника", max(min_score, 65)
        if cands:
            ranked = sorted(((name_score(m.norm, sup[sid].norm), sid) for sid in cands), key=lambda x: (-x[0], x[1]))
            best, sid = ranked[0]
            gap = best - ranked[1][0] if len(ranked) > 1 else 100
            clash = pair_conflict(m, sup[sid])
            if best >= min_score and gap >= 3 and not clash:
                pairs.append(Pair(m, sup[sid], method, best, sup_id))
                continue
            if clash:
                reason = f"{method} збігся, але {clash}"
            elif best < min_score:
                reason = f"{method} збігся, але назви різні"
            else:
                reason = f"{method} збігся з кількома товарами"
            fallback = (sid, best, reason)

        # 2. однакова назва
        same = [sid for sid in name_idx.get(m.norm, ()) if (m.id, sid) not in blocked]
        if len(same) == 1 and not pair_conflict(m, sup[same[0]]):
            pairs.append(Pair(m, sup[same[0]], "назва", 100.0, sup_id))
            continue

        # 3. та сама модель того самого бренду, а в назві постачальника лише більше опису
        strong = strong_ident(model_ident(m))
        if strong:
            ids = set.intersection(*(ident_idx.get(t, set()) for t in strong))
            cands = [sid for sid in ids if (m.id, sid) not in blocked and model_match_ok(m, sup[sid])]
            # підозрілі сусіди: ті самі позначення моделі, але товар не пройшов перевірку
            rivals = [sid for sid in ids if sid not in cands and not is_bundle(sup[sid]) and (m.id, sid) not in blocked]
            if len(cands) == 1 and not rivals:
                sc = name_score(m.norm, sup[cands[0]].norm)
                if sc >= 70:
                    pairs.append(Pair(m, sup[cands[0]], "модель", sc, sup_id))
                    continue

        # 4. схожа назва
        counter = Counter()
        for t in set(m.norm.split()):
            ids = tok_idx.get(t)
            if ids and len(ids) <= df_limit:
                counter.update(ids)
        best_sid, best, second = None, 0.0, 0.0
        for sid, _ in counter.most_common(300):
            if (m.id, sid) in blocked:
                continue
            sc = name_score(m.norm, sup[sid].norm)
            if sc > best:
                best_sid, best, second = sid, sc, best
            elif sc > second:
                second = sc
        if best_sid:
            s = sup[best_sid]
            clash = pair_conflict(m, s)
            # слово чи бренд, що є лише в одній назві, часто означає інше виконання товару
            extra = (m.words ^ s.words) | brand_diff(m.brands, s.brands)
            if (best >= cfg["fuzzy_auto_accept"] and m.model and m.model == s.model
                    and not clash and not extra and best - second >= 2):
                pairs.append(Pair(m, s, "схожа назва", best, sup_id))
                continue
            if best >= cfg["fuzzy_suggest"] and (fallback is None or best > fallback[1]):
                if clash:
                    reason = f"назви схожі, але {clash}"
                elif extra:
                    reason = "назви схожі, різняться словами: " + ", ".join(sorted(extra)[:3]).lower()
                    if len(extra) > 3:
                        reason += f" та ще {len(extra) - 3}"
                elif len(same) > 1:
                    reason = "кілька товарів з такою назвою"
                else:
                    reason = "назви схожі, перевірте"
                fallback = (best_sid, best, reason)

        if fallback:
            sid, sc, reason = fallback
            suggestions.append({"sup": sup_id, "mid": m.id, "sid": sid, "score": sc, "reason": reason})

    suggestions.sort(key=lambda x: -x["score"])
    return pairs, suggestions, warnings


# ---------------------------------------------------------------- історія

def classify(mp, sp, cfg) -> str:
    if mp is None or sp is None:
        return "noprice"
    lo = sp * (1 + cfg["min_markup_pct"] / 100) - cfg["tolerance_uah"]
    hi = sp * (1 + cfg["max_markup_pct"] / 100) + cfg["tolerance_uah"]
    if mp < lo:
        return "below"
    if mp > hi:
        return "above"
    return "ok"


SNAP_HEADER = ["supplier", "my_id", "supplier_id", "my_price", "supplier_price", "my_available", "supplier_available"]


def load_history(today: date, days: int, default: str) -> tuple[dict, list[date]]:
    files = []
    for path in SNAP_DIR.glob("*.csv"):
        try:
            d = date.fromisoformat(path.stem)
        except ValueError:
            continue
        if d < today:
            files.append((d, path))
    files.sort()
    start = today - timedelta(days=days - 1)
    chosen = [(d, p) for d, p in files if d >= start]
    if files and (not chosen or chosen[-1][0] != files[-1][0]):
        chosen.append(files[-1])
    hist = defaultdict(list)
    for d, path in sorted(chosen):
        with path.open(encoding="utf-8") as f:
            for r in csv.DictReader(f):
                key = (r.get("supplier") or default, r["my_id"], r["supplier_id"])
                hist[key].append({
                    "d": d,
                    "mp": parse_price(r.get("my_price")),
                    "sp": parse_price(r.get("supplier_price")),
                    "sa": r.get("supplier_available") == "1",
                    "ma": r.get("my_available", "1") == "1",
                })
    return hist, [d for d, _ in sorted(chosen)]


def write_snapshot(today: date, pairs: list[Pair]) -> None:
    SNAP_DIR.mkdir(parents=True, exist_ok=True)
    with (SNAP_DIR / f"{today.isoformat()}.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(SNAP_HEADER)
        for p in sorted(pairs, key=lambda p: p.key):
            w.writerow([p.sup, p.m.id, p.s.id, fmt_num(p.m.price), fmt_num(p.s.price),
                        int(p.m.available), int(p.s.available)])


CHANGES_HEADER = ["date", "supplier", "side", "my_id", "supplier_id", "name", "old_price", "new_price", "delta", "delta_pct"]


def update_changes(today: date, pairs: list[Pair], hist: dict, default: str) -> list[dict]:
    events = []
    for p in pairs:
        series = hist.get(p.key)
        if not series:
            continue
        prev = series[-1]
        for side, old, new in ((SIDE_SUP, prev["sp"], p.s.price), (SIDE_MY, prev["mp"], p.m.price)):
            if old is not None and new is not None and abs(new - old) >= 0.01:
                events.append({"date": today.isoformat(), "supplier": p.sup, "side": side, "my_id": p.m.id,
                               "supplier_id": p.s.id, "name": p.m.name, "old_price": fmt_num(old),
                               "new_price": fmt_num(new), "delta": fmt_num(round(new - old, 2)),
                               "delta_pct": pct(new, old)})
    rows = []
    if CHANGES_CSV.exists():
        with CHANGES_CSV.open(encoding="utf-8") as f:
            for r in csv.DictReader(f):
                if r["date"] == today.isoformat():
                    continue
                r.setdefault("supplier", default)
                r["supplier"] = r.get("supplier") or default
                rows.append({k: r.get(k, "") for k in CHANGES_HEADER})
    rows += events
    rows.sort(key=lambda r: (r["date"], r["supplier"], r["side"], r["my_id"]))
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    with CHANGES_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=CHANGES_HEADER, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)
    return rows


# ---------------------------------------------------------------- рядки звіту

def pick_reference(cells: list[dict]) -> dict | None:
    """Орієнтир для порівняння: найдешевший постачальник, у якого товар є в наявності."""
    priced = [c for c in cells if c["sp"] is not None]
    if not priced:
        return None
    live = [c for c in priced if c["savail"]]
    return min(live or priced, key=lambda c: c["sp"])


def build_rows(pairs, hist, changes, today, cfg, suppliers) -> list[dict]:
    by_key = defaultdict(list)
    for e in changes:
        by_key[(e["supplier"], e["my_id"], e["supplier_id"])].append(e)
    month_ago = (today - timedelta(days=29)).isoformat()
    order_sup = {sp["id"]: i for i, sp in enumerate(suppliers)}

    grouped = defaultdict(list)
    for p in pairs:
        grouped[p.m.id].append(p)

    rows = []
    for mid, plist in grouped.items():
        m = plist[0].m
        mp = m.price
        cells = []
        for p in sorted(plist, key=lambda p: order_sup.get(p.sup, 99)):
            series = hist.get(p.key, [])
            prev = series[-1] if series else None
            sp = p.s.price
            sch = round(sp - prev["sp"], 2) if prev and prev["sp"] is not None and sp is not None and abs(sp - prev["sp"]) >= 0.01 else 0
            sup_events = [e for e in by_key.get(p.key, []) if e["side"] == SIDE_SUP]
            last = sup_events[-1] if sup_events else None
            cells.append({
                "sup": p.sup, "sid": p.s.id, "sname": p.s.name, "surl": p.s.url,
                "sp": sp, "savail": p.s.available, "sch": sch, "how": p.how, "score": p.score,
                "diff": round(mp - sp, 2) if mp is not None and sp is not None else None,
                "diffp": pct(mp, sp) if mp is not None and sp is not None else None,
                "c30": sum(1 for e in sup_events if e["date"] >= month_ago),
                "last": last["date"] if last else "",
                "series": series, "key": list(p.key),
            })

        ref = pick_reference(cells)
        sp = ref["sp"] if ref else None
        sch = ref["sch"] if ref else 0
        prev_mine = None
        if ref:
            ser = ref["series"]
            prev_mine = ser[-1]["mp"] if ser else None
        mch = round(mp - prev_mine, 2) if mp is not None and prev_mine is not None and abs(mp - prev_mine) >= 0.01 else 0

        st = classify(mp, sp, cfg)
        nostock = bool(cells) and not any(c["savail"] for c in cells)
        attention = st == "below" or (sch != 0 and mch == 0)
        sev = "red" if attention else ("orange" if st == "above" else "ok")
        diff = round(mp - sp, 2) if mp is not None and sp is not None else None
        diffp = pct(mp, sp) if mp is not None and sp is not None else None

        what = []
        if sch:
            old = sp - sch
            what.append(f"Постачальник змінив ціну: {fmt_num(old)} → {fmt_num(sp)} грн ({pct(sp, old):+.1f}%)"
                        + (", ваша ціна не змінилась" if not mch else ", ваша теж змінилась"))
        if st == "below":
            what.append(f"Ваша ціна нижча на {fmt_num(abs(diff))} грн ({abs(diffp):.1f}%)")
        elif st == "above":
            what.append(f"Ваша ціна вища на {fmt_num(diff)} грн ({diffp:.1f}%)")
        if nostock and m.available:
            what.append("Немає в наявності в жодного постачальника")

        days = 0
        pts = []
        if ref:
            full = ref["series"] + [{"d": today, "mp": mp, "sp": sp}]
            start_d = None
            for e in reversed(full):
                if classify(e["mp"], e["sp"], cfg) in ("ok", "noprice"):
                    break
                start_d = e["d"]
            days = (today - start_d).days + 1 if start_d else 0
            lastv = None
            for e in full:
                v = (e["mp"], e["sp"])
                if v != lastv:
                    pts.append([e["d"].isoformat(), e["mp"], e["sp"]])
                    lastv = v
            if pts and pts[-1][0] != today.isoformat():
                pts.append([today.isoformat(), mp, sp])

        for c in cells:
            c.pop("series", None)

        rows.append({
            "mid": m.id, "mname": m.name, "murl": m.url, "mp": mp, "mavail": m.available,
            "cells": cells, "ref": ref["sup"] if ref else "", "sid": ref["sid"] if ref else "",
            "sname": ref["sname"] if ref else "", "surl": ref["surl"] if ref else "",
            "sp": sp, "savail": ref["savail"] if ref else False, "diff": diff, "diffp": diffp,
            "st": st, "sev": sev, "sch": sch, "mch": mch, "nostock": nostock, "what": "; ".join(what),
            "days": days, "c30": ref["c30"] if ref else 0, "last": ref["last"] if ref else "",
            "how": ref["how"] if ref else "", "score": ref["score"] if ref else 0, "hist": pts,
        })

    order = {"red": 0, "orange": 1, "ok": 2}
    rows.sort(key=lambda r: (order[r["sev"]], 0 if r["sch"] else 1, -abs(r["diffp"] or 0), r["mname"]))
    return rows


# ---------------------------------------------------------------- звіти

def site_repo() -> str:
    raw = os.environ.get("SITE_REPO", "").strip()
    raw = re.sub(r"^https?://github\.com/", "", raw).strip("/").removesuffix(".git").strip("/")
    return raw if re.fullmatch(r"[\w.-]+/[\w.-]+", raw) else ""


def pages_url(repo: str) -> str:
    owner, name = repo.split("/", 1)
    if name.lower() == f"{owner.lower()}.github.io":
        return f"https://{name.lower()}/"
    return f"https://{owner.lower()}.github.io/{name}/"


def repo_links(cfg: dict) -> dict:
    repo = os.environ.get("GITHUB_REPOSITORY", "")
    branch = os.environ.get("GITHUB_REF_NAME", "main")
    site = pages_url(site_repo()) if site_repo() and os.environ.get("DASHBOARD_PASSWORD") else ""
    if not repo:
        return {"repo": "", "pages": site, "edit_sugg": "", "edit_map": "", "branch": branch}
    return {
        "repo": f"https://github.com/{repo}",
        "pages": site or (pages_url(repo) if cfg.get("github_pages") else ""),
        "edit_sugg": f"https://github.com/{repo}/edit/{branch}/mapping/suggestions.csv",
        "edit_map": f"https://github.com/{repo}/edit/{branch}/mapping/mapping.csv",
        "branch": branch,
    }


def write_suggestions(suggestions, mine, catalogs) -> None:
    MAP_DIR.mkdir(parents=True, exist_ok=True)
    with SUGG_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(SUGG_HEADER)
        for s in suggestions:
            m, o = mine[s["mid"]], catalogs[s["sup"]][s["sid"]]
            w.writerow(["", s["sup"], fmt_num(s["score"]), m.id, o.id, m.name, o.name,
                        fmt_num(m.price), fmt_num(o.price), s["reason"]])


def write_html(ctx: dict) -> None:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    (DOCS_DIR / ".nojekyll").write_text("", encoding="utf-8")
    data = json.dumps(ctx, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    html = TEMPLATE.read_text(encoding="utf-8").replace("/*__DATA__*/null", data)
    (DOCS_DIR / "index.html").write_text(html, encoding="utf-8")


def write_site(ctx: dict, cfg: dict) -> None:
    """Зашифрована копія дашборда для публікації за посиланням. Без пароля не збирається."""
    password = os.environ.get("DASHBOARD_PASSWORD", "")
    if not password:
        return
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM

    b64 = lambda b: base64.b64encode(b).decode("ascii")  # noqa: E731
    # сіль стала для репозиторію, тому «запам'ятати на пристрої» не злітає після щоденного оновлення
    salt = hashlib.sha256(f"price-watch|{site_repo()}".encode("utf-8")).digest()[:16]
    key = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, PBKDF2_ITER, dklen=32)
    iv = os.urandom(12)
    packed = gzip.compress((DOCS_DIR / "index.html").read_bytes(), compresslevel=9, mtime=0)
    payload = {"salt": b64(salt), "iv": b64(iv), "iter": PBKDF2_ITER, "ct": b64(AESGCM(key).encrypt(iv, packed, None))}
    names = ", ".join(sp["label"] for sp in ctx["suppliers"])
    title = htmllib.escape(f"Ціни: {cfg['my_label']} і {names}"[:90])
    page = (SITE_TEMPLATE.read_text(encoding="utf-8")
            .replace("__TITLE__", title)
            .replace("__UPDATED__", htmllib.escape(ctx["generated"]))
            .replace("/*__PAYLOAD__*/null", json.dumps(payload)))
    SITE_DIR.mkdir(parents=True, exist_ok=True)
    (SITE_DIR / "index.html").write_text(page, encoding="utf-8")
    (SITE_DIR / ".nojekyll").write_text("", encoding="utf-8")


def md_cell(text) -> str:
    return str(text if text is not None else "").replace("|", "\\|").replace("\n", " ")


def sup_label(ctx: dict, sid: str) -> str:
    for sp in ctx["suppliers"]:
        if sp["id"] == sid:
            return sp["label"]
    return sid


def build_readme(ctx: dict) -> str:
    c, labels, links = ctx["counts"], ctx["labels"], ctx["links"]
    rows = ctx["rows"]
    sup_names = ", ".join(sp["label"] for sp in ctx["suppliers"])
    out = [f"# Ціни: {labels['my']} і постачальники", "", f"Оновлено {ctx['generated']} за Києвом.", ""]
    if links["pages"]:
        out += [f"**[Відкрити веб-версію]({links['pages']})**", ""]
    for w in ctx["warnings"]:
        out.append(f"> ⚠️ {w}  ")
    if ctx["warnings"]:
        out.append("")
    out += [
        f"🔴 Потребують уваги: **{c['attention']}**  ",
        f"🟠 Ваша ціна вища за орієнтир: **{c['above']}**  ",
        f"🟢 Ціни в межах норми: **{c['ok']}**  ",
        f"📦 Немає в наявності в жодного постачальника: **{c['nostock']}**  ",
        f"Спільних товарів: {c['matched']} (у вашому фіді {ctx['stats']['my_total']})",
        f"Постачальники: {sup_names}",
        "",
    ]
    report_links = ["[report.xlsx](report.xlsx)", "[журнал змін](data/changes.csv)",
                     f"[пари на перевірку ({len(ctx['sugg'])})](mapping/suggestions.csv)"]
    out += ["Повний звіт: " + ", ".join(report_links), ""]

    changed = [r for r in rows if r["sch"]]
    if changed:
        out += ["## 🔴 Постачальник змінив ціну", "", "| Товар | Постачальник | Було | Стало | Зміна | Ваша ціна |",
                "|---|---|---:|---:|---:|---:|"]
        for r in changed[:60]:
            old = r["sp"] - r["sch"]
            out.append(f"| {md_cell(r['mname'])} | {md_cell(sup_label(ctx, r['ref']))} | {fmt_num(old)} | {fmt_num(r['sp'])} "
                       f"| {pct(r['sp'], old):+.1f}% | {fmt_num(r['mp'])} |")
        out.append("")
    below = [r for r in rows if r["st"] == "below" and not r["sch"]]
    if below:
        out += ["## 🔴 Ваша ціна нижча за постачальника", "", "| Товар | Ваша | Постачальник | Ціна | Різниця | Днів |",
                "|---|---:|---|---:|---:|---:|"]
        for r in below[:60]:
            out.append(f"| {md_cell(r['mname'])} | {fmt_num(r['mp'])} | {md_cell(sup_label(ctx, r['ref']))} "
                       f"| {fmt_num(r['sp'])} | {r['diffp']:+.1f}% | {r['days']} |")
        out.append("")
    above = [r for r in rows if r["st"] == "above" and not r["sch"]]
    if above:
        out += ["## 🟠 Ваша ціна вища за постачальника", "", "| Товар | Ваша | Постачальник | Ціна | Різниця | Днів |",
                "|---|---:|---|---:|---:|---:|"]
        for r in above[:40]:
            out.append(f"| {md_cell(r['mname'])} | {fmt_num(r['mp'])} | {md_cell(sup_label(ctx, r['ref']))} "
                       f"| {fmt_num(r['sp'])} | {r['diffp']:+.1f}% | {r['days']} |")
        out.append("")
    gone = [r for r in rows if r["nostock"] and r["mavail"]]
    if gone:
        out += ["## 📦 Немає в наявності в жодного постачальника", "", "| Товар | Ваша ціна | Постачальники |", "|---|---:|---|"]
        for r in gone[:40]:
            out.append(f"| {md_cell(r['mname'])} | {fmt_num(r['mp'])} | "
                       f"{md_cell(', '.join(sup_label(ctx, x['sup']) for x in r['cells']))} |")
        out.append("")
    out += ["---", f"Як це працює і як налаштувати: [ІНСТРУКЦІЯ.md]({quote('ІНСТРУКЦІЯ.md')})", ""]
    return "\n".join(out)


def telegram_text(ctx: dict) -> str:
    c, labels = ctx["counts"], ctx["labels"]
    lines = [f"Ціни {labels['my']}, {ctx['generated']}",
             f"🔴 Потребують уваги: {c['attention']} (постачальник змінив ціну: {c['changed']}, ваша нижча: {c['below']})",
             f"🟠 Ваша вища: {c['above']}"]
    if c["nostock"]:
        lines.append(f"📦 Немає в наявності в жодного постачальника: {c['nostock']}")
    changed = [r for r in ctx["rows"] if r["sch"]][:12]
    if changed:
        lines.append("")
        for r in changed:
            old = r["sp"] - r["sch"]
            lines.append(f"• {r['mname'][:60]} [{sup_label(ctx, r['ref'])}]: {fmt_num(old)} → {fmt_num(r['sp'])} "
                         f"({pct(r['sp'], old):+.1f}%), у вас {fmt_num(r['mp'])}")
    for w in ctx["warnings"]:
        lines.append(f"⚠️ {w}")
    if ctx["links"]["pages"]:
        lines += ["", ctx["links"]["pages"]]
    return "\n".join(lines)[:4000]


def send_telegram(ctx: dict, cfg: dict) -> None:
    token, chat = os.environ.get("TELEGRAM_BOT_TOKEN", ""), os.environ.get("TELEGRAM_CHAT_ID", "")
    if not token or not chat or os.environ.get("SKIP_TELEGRAM"):
        return
    if cfg.get("telegram_only_if_attention") and not ctx["counts"]["attention"] and not ctx["warnings"]:
        return
    api = f"https://api.telegram.org/bot{token}"
    try:
        r = requests.post(f"{api}/sendMessage", timeout=30,
                          data={"chat_id": chat, "text": telegram_text(ctx), "disable_web_page_preview": "true"})
        if r.status_code != 200:
            print(f"Telegram відповів {r.status_code}: {r.text[:200]}")
            return
        page = DOCS_DIR / "index.html"
        if cfg.get("telegram_attach_dashboard") and not ctx["links"]["pages"] and page.exists():
            fname = f"prices-{ctx['today']}.html"
            r = requests.post(f"{api}/sendDocument", timeout=60,
                              data={"chat_id": chat, "caption": "Дашборд: відкрийте файл у браузері"},
                              files={"document": (fname, page.read_bytes(), "text/html")})
            if r.status_code != 200:
                print(f"Telegram не прийняв файл дашборда: {r.status_code} {r.text[:200]}")
    except Exception as exc:  # noqa: BLE001
        print("Telegram недоступний:", str(exc).replace(token, "***"))


# ---------------------------------------------------------------- головне

def main() -> None:
    cfg = load_config()
    now = datetime.now(KYIV)
    today = date.fromisoformat(os.environ["RUN_DATE"]) if os.environ.get("RUN_DATE") else now.date()
    my_src = os.environ.get("MY_FEED_URL", "").strip()
    suppliers = load_suppliers(cfg)
    default_sup = suppliers[0]["id"]

    absorbed = absorb_decisions(today, default_sup)
    if absorbed:
        print(f"Перенесено рішень із suggestions.csv у mapping.csv: {absorbed}")

    print("Завантажую джерела...")
    try:
        mine = parse_source(load_feed(my_src, cfg["my_label"]), cfg["my_label"], my_src, {})
    except FeedError as exc:
        sys.exit(str(exc))
    print(f"  {cfg['my_label']}: {len(mine)} товарів")
    if not mine:
        sys.exit("Ваш фід не містить жодного товару, звіт не оновлюю.")

    catalogs, warnings = {}, []
    for sp in suppliers:
        if not sp["src"] and not sp["scraper"]:
            warnings.append(f"Для постачальника «{sp['label']}» не вказано ні посилання, ні файл.")
            continue
        try:
            if sp["scraper"]:
                cat = run_scraper(sp)
            else:
                cat = parse_source(load_feed(sp["src"], sp["label"]), sp["label"], sp["src"], sp["columns"], sp["format"])
        except FeedError as exc:
            warnings.append(f"{exc} Звіт побудовано без цього постачальника.")
            print("  Увага:", exc)
            continue
        if not cat:
            warnings.append(f"У джерелі «{sp['label']}» немає жодного товару, цього разу пропускаю.")
            continue
        catalogs[sp["id"]] = cat
        print(f"  {sp['label']}: {len(cat)} товарів")
    if not catalogs:
        sys.exit("Жоден фід постачальника не прочитався, звіт не оновлюю.")
    suppliers = [sp for sp in suppliers if sp["id"] in catalogs]

    forced, blocked, excluded = load_rules(suppliers)
    pairs, suggestions = [], []
    for sp in suppliers:
        sid = sp["id"]
        p, sg, w = match_offers(mine, catalogs[sid], cfg, forced[sid], blocked[sid], excluded[sid], sid)
        pairs += p
        suggestions += sg
        warnings += [f"{sp['label']}: {x}" for x in w]
    suggestions.sort(key=lambda x: -x["score"])

    priced = [p for p in pairs if p.m.price and p.s.price]
    if len(priced) < len(pairs):
        warnings.append(f"Без ціни в одному з джерел: {len(pairs) - len(priced)} пар, їх не порівнюю.")

    last = json.loads(LAST_RUN.read_text(encoding="utf-8")) if LAST_RUN.exists() else {}
    if last and last.get("date") != today.isoformat():
        if last.get("matched", 0) >= 20 and len(priced) < last["matched"] * 0.5:
            warnings.append(f"Спільних позицій стало {len(priced)} замість {last['matched']}. Перевірте, чи фіди віддаються повністю.")
        if last.get("my_total", 0) >= 50 and len(mine) < last["my_total"] * 0.6:
            warnings.append(f"У вашому фіді {len(mine)} товарів, минулого разу було {last['my_total']}.")
        for sp in suppliers:
            was = (last.get("sup_totals") or {}).get(sp["id"], 0)
            if was >= 50 and len(catalogs[sp["id"]]) < was * 0.6:
                warnings.append(f"У «{sp['label']}» {len(catalogs[sp['id']])} товарів, минулого разу було {was}.")

    hist, hist_dates = load_history(today, cfg["report_history_days"], default_sup)
    changes = update_changes(today, priced, hist, default_sup)
    write_snapshot(today, priced)
    rows = build_rows(priced, hist, changes, today, cfg, suppliers)
    write_suggestions(suggestions, mine, catalogs)

    labels = {sp["id"]: sp["label"] for sp in suppliers}
    log_from = (today - timedelta(days=cfg["log_days"] - 1)).isoformat()
    log = [{"d": e["date"], "side": e["side"], "sup": e["supplier"], "supLabel": labels.get(e["supplier"], e["supplier"]),
            "mid": e["my_id"], "sid": e["supplier_id"], "name": e["name"],
            "old": parse_price(e["old_price"]), "new": parse_price(e["new_price"]),
            "pct": float(e["delta_pct"]) if e["delta_pct"] not in ("", None) else None}
           for e in reversed(changes) if e["date"] >= log_from and e["supplier"] in labels]

    ctx = {
        "generated": (now if not os.environ.get("RUN_DATE") else datetime.combine(today, now.timetz())).strftime("%d.%m.%Y %H:%M"),
        "today": today.isoformat(),
        "historyDays": cfg["report_history_days"],
        "logDays": cfg["log_days"],
        "labels": {"my": cfg["my_label"], "sup": suppliers[0]["label"]},
        "suppliers": [{"id": sp["id"], "label": sp["label"]} for sp in suppliers],
        "thresholds": {k: cfg[k] for k in ("tolerance_uah", "min_markup_pct", "max_markup_pct")},
        "counts": {
            "matched": len(rows),
            "pairs": len(priced),
            "attention": sum(r["sev"] == "red" for r in rows),
            "changed": sum(bool(r["sch"]) for r in rows),
            "below": sum(r["st"] == "below" for r in rows),
            "above": sum(r["st"] == "above" for r in rows),
            "ok": sum(r["st"] == "ok" for r in rows),
            "nostock": sum(r["nostock"] and r["mavail"] for r in rows),
            "myout": sum(not r["mavail"] and any(c["savail"] for c in r["cells"]) for r in rows),
        },
        "stats": {"my_total": len(mine), "sup_total": sum(len(c) for c in catalogs.values()),
                  "sup_totals": {sid: len(cat) for sid, cat in catalogs.items()},
                  "by_method": dict(Counter(p.how for p in priced))},
        "rows": rows,
        "log": log,
        "sugg": [{"sup": x["sup"], "supLabel": labels.get(x["sup"], x["sup"]), "mid": x["mid"], "sid": x["sid"],
                  "mname": mine[x["mid"]].name, "sname": catalogs[x["sup"]][x["sid"]].name,
                  "mp": mine[x["mid"]].price, "sp": catalogs[x["sup"]][x["sid"]].price,
                  "score": x["score"], "reason": x["reason"],
                  "murl": mine[x["mid"]].url, "surl": catalogs[x["sup"]][x["sid"]].url} for x in suggestions],
        "warnings": warnings,
        "links": repo_links(cfg),
    }

    write_html(ctx)
    write_site(ctx, cfg)
    from report_xlsx import write_xlsx  # локальний модуль
    write_xlsx(ROOT / "report.xlsx", ctx, hist, hist_dates, today)
    readme = build_readme(ctx)
    (ROOT / "README.md").write_text(readme, encoding="utf-8")
    LAST_RUN.write_text(json.dumps({"date": today.isoformat(), "matched": len(priced), "my_total": len(mine),
                                    "sup_totals": ctx["stats"]["sup_totals"]}, ensure_ascii=False, indent=2),
                        encoding="utf-8")

    summary_path = os.environ.get("GITHUB_STEP_SUMMARY")
    if summary_path:
        with open(summary_path, "a", encoding="utf-8") as f:
            f.write(readme)
    send_telegram(ctx, cfg)

    c = ctx["counts"]
    print(f"Спільних товарів: {c['matched']}, пар з постачальниками: {c['pairs']} "
          f"({', '.join(f'{k}: {v}' for k, v in ctx['stats']['by_method'].items())})")
    print(f"Потребують уваги: {c['attention']}, змінив ціну постачальник: {c['changed']}, "
          f"ваша нижча: {c['below']}, ваша вища: {c['above']}, немає в наявності ніде: {c['nostock']}, "
          f"пар на перевірку: {len(suggestions)}")
    for w in warnings:
        print("Увага:", w)


if __name__ == "__main__":
    main()
